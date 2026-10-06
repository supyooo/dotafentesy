from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

F = "Arial"
def f(**k): return Font(name=F, **k)
BLUE = f(color="0000FF"); GREEN = f(color="008000"); BOLD = f(bold=True); NORM = f(); GREY = f(color="666666", size=9)
YEL = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="E8EEF7"); MERGE = PatternFill("solid", fgColor="FFF4CC")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(bottom=thin)
leagues = ["Herald","Guardian","Crusader","Archon","Legend","Ancient","Divine","Immortal"]
active = [30,45,60,40,25,12,7,3]
target = [.10,.15,.20,.20,.15,.10,.07,.03]
calib  = [.30,.20,.20,.15,.15,0,0,0]
wb = Workbook()

# ---------- Ввод ----------
s = wb.active; s.title = "Ввод"
s["A1"] = "Модель лиг — ввод"; s["A1"].font = f(bold=True, size=14)
s["A2"] = "Меняй только жёлтые ячейки с синим текстом. Остальные листы пересчитаются сами."; s["A2"].font = GREY
s["A3"] = "Пример заполнения: внизу лестницы людей много, в Divine 7, в Immortal 3 — случай «верхние лиги почти пустые»."; s["A3"].font = GREY
hdr = ["Лига","Активных игроков","Целевая доля","Доля новичков при калибровке"]
for j,h in enumerate(hdr,1):
    c = s.cell(4,j,h); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
for i,lg in enumerate(leagues):
    r = 5+i
    s.cell(r,1,lg).font = NORM
    for j,(v,fmt) in enumerate([(active[i],"0"),(target[i],"0%"),(calib[i],"0%")],2):
        c = s.cell(r,j,v); c.font = BLUE; c.fill = YEL; c.number_format = fmt
s["A13"] = "Итого"; s["A13"].font = BOLD
for col,fmt in [("B","0"),("C","0%"),("D","0%")]:
    s[f"{col}13"] = f"=SUM({col}5:{col}12)"; s[f"{col}13"].font = BOLD; s[f"{col}13"].number_format = fmt
s["E13"] = '=IF(OR(ROUND(C13,4)<>1,ROUND(D13,4)<>1),"Доли должны давать в сумме 100%","")'; s["E13"].font = f(color="C00000")
s["C4"].comment = Comment("Целевое распределение игроков по лигам. Плавающие зоны тянут лестницу к нему. Значения — предложение, не данные.", "model")
s["D4"].comment = Comment("Из таблицы калибровки в документе: ниже 30-го процентиля — Herald, 30–49 — Guardian, 50–69 — Crusader, 70–84 — Archon, 85+ — Legend.", "model")

params = [
 ("Размер группы (цель)",30,"0","Группы делятся поровну, число групп = округлить(игроков / 30)."),
 ("Минимум игроков, чтобы лига играла своими группами",20,"0","Меньше — лига играет в сводной таблице с соседней лигой ниже."),
 ("Базовая зона повышения / вылета",.20,"0%","Зона при точном попадании лиги в целевую долю."),
 ("Минимальная зона",.10,"0%",""),
 ("Максимальная зона",.30,"0%",""),
 ("Чувствительность: сдвиг зоны при 100% недоборе",.10,"0%","Лига пустая (недобор 100%) — зона в неё шире на это значение. Переполнена вдвое — уже на столько же."),
 ("Immortal: минимум игроков для вылета",30,"0","Меньше — из Immortal никто не вылетает."),
 ("Новичков за раунд (для прогноза)",50,"0","Распределяются по лигам по калибровке."),
 ("Отток за раунд (для прогноза)",.10,"0%","Доля игроков каждой лиги, которые не вернулись."),
]
s["A15"] = "Параметры"; s["A15"].font = f(bold=True, size=12)
for i,(lab,v,fmt,note) in enumerate(params):
    r = 16+i
    s.cell(r,1,lab).font = NORM
    c = s.cell(r,2,v); c.font = BLUE; c.fill = YEL; c.number_format = fmt
    s.cell(r,3,note).font = GREY
s.column_dimensions["A"].width = 52; s.column_dimensions["B"].width = 18; s.column_dimensions["C"].width = 16; s.column_dimensions["D"].width = 18; s.column_dimensions["E"].width = 36
s.row_dimensions[4].height = 30

P = lambda r: f"'Ввод'!$B${r}"
SIZE,MINL,BASE,ZMIN,ZMAX,K,IMM,NEWC,CHURN = (P(r) for r in range(16,25))
def clamp(x): return f"MIN({ZMAX},MAX({ZMIN},{x}))"

# ---------- Группы ----------
g = wb.create_sheet("Группы")
g["A1"] = "Раунд по введённым числам: таблицы, группы, зоны"; g["A1"].font = f(bold=True, size=14)
g["A2"] = "Жёлтые строки — лиги в сводной таблице (в лиге меньше минимума). Перемещения — ожидаемые: считаем, что игроки каждой лиги распределены по таблице равномерно."; g["A2"].font = GREY
cols = ["Лига","Активных","Доля факт","Цель","Недобор","Накопл. (служ.)","Пул (служ.)","Таблица №","Играет вместе с","Игроков в таблице","Групп","Размеры групп","Зона повышения","Зона вылета","В каждой группе","Поднимутся (ожид.)","Опустятся (ожид.)","В следующем раунде"]
for j,h in enumerate(cols,1):
    c = g.cell(4,j,h); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
for i in range(8):
    r = 5+i
    g[f"A{r}"] = f"='Ввод'!A{r}"; g[f"A{r}"].font = GREEN
    g[f"B{r}"] = f"='Ввод'!B{r}"; g[f"B{r}"].font = GREEN
    g[f"C{r}"] = f"=IF(SUM($B$5:$B$12)=0,0,B{r}/SUM($B$5:$B$12))"
    g[f"D{r}"] = f"='Ввод'!C{r}"; g[f"D{r}"].font = GREEN
    g[f"E{r}"] = f"=IF(D{r}=0,0,(D{r}-C{r})/D{r})"
    if i < 7:
        if i == 0:
            g["F5"] = "=B5"; g["G5"] = 1
        else:
            g[f"F{r}"] = f"=IF(F{r-1}>={MINL},B{r},F{r-1}+B{r})"
            g[f"G{r}"] = f"=IF(F{r-1}>={MINL},G{r-1}+1,G{r-1})"
        g[f"H{r}"] = f"=IF(AND(G{r}=$G$11,$F$11<{MINL},$G$11>1),$G$11-1,G{r})"
        parts = ",".join(f'IF($H${5+k}=H{r},$A${5+k},"")' for k in range(7))
        g[f"I{r}"] = f'=IF(COUNTIF($H$5:$H$11,H{r})=1,"своя таблица",_xlfn.TEXTJOIN(" + ",TRUE,{parts}))'
        g[f"J{r}"] = f"=SUMIF($H$5:$H$11,H{r},$B$5:$B$11)"
        g[f"K{r}"] = f"=IF(J{r}=0,0,MAX(1,ROUND(J{r}/{SIZE},0)))"
        g[f"L{r}"] = f'=IF(K{r}=0,"—",IF(MOD(J{r},K{r})=0,K{r}&" × "&J{r}/K{r},MOD(J{r},K{r})&" × "&(INT(J{r}/K{r})+1)&" + "&(K{r}-MOD(J{r},K{r}))&" × "&INT(J{r}/K{r})))'
        g[f"M{r}"] = f"={clamp(f'{BASE}+{K}*E{r+1}')}"
        g[f"N{r}"] = "=0" if i == 0 else f"={clamp(f'{BASE}-{K}*E{r}')}"
        g[f"O{r}"] = (f'=IF(K{r}=0,"—","вверх "&ROUNDDOWN(SUMPRODUCT(($H$5:$H$11=H{r})*$B$5:$B$11*$M$5:$M$11)/K{r},0)'
                      f'&" / вниз "&ROUNDDOWN(SUMPRODUCT(($H$5:$H$11=H{r})*$B$5:$B$11*$N$5:$N$11)/K{r},0))')
    else:
        g["H12"] = "—"; g["I12"] = "одна общая таблица"; g["J12"] = "=B12"; g["K12"] = "=IF(B12>0,1,0)"
        g["L12"] = '=IF(B12>0,B12&" (без групп)","—")'
        g["M12"] = "=0"
        g["N12"] = f"=IF(B12<{IMM},0,{clamp(f'{BASE}-{K}*E12')})"
        g["O12"] = '="вверх 0 / вниз "&ROUNDDOWN(B12*N12,0)'
    g[f"P{r}"] = f"=B{r}*M{r}"
    g[f"Q{r}"] = f"=B{r}*N{r}"
    inn = (f"+P{r-1}" if i>0 else "") + (f"+Q{r+1}" if i<7 else "")
    g[f"R{r}"] = f"=B{r}-P{r}-Q{r}{inn}"
g["A13"] = "Итого"; g["A13"].font = BOLD
for col in "BCJPQR":
    g[f"{col}13"] = f"=SUM({col}5:{col}12)" if col!="J" else "=SUM(B5:B12)"
    g[f"{col}13"].font = BOLD
for r in range(5,14):
    for col,fmt in {"B":"0","C":"0.0%","D":"0%","E":"0%;-0%;0%","F":"0","G":"0","J":"0","K":"0","M":"0%","N":"0%","P":"0.0","Q":"0.0","R":"0.0"}.items():
        g[f"{col}{r}"].number_format = fmt
    for col in "ABCDEFGHIJKLMNOPQR":
        if g[f"{col}{r}"].font.color is None or g[f"{col}{r}"].font != GREEN:
            if g[f"{col}{r}"].font != GREEN and r<13: g[f"{col}{r}"].font = NORM
for col in "FG":
    for r in range(4,13): g[f"{col}{r}"].font = GREY
g.conditional_formatting.add("A5:R11", FormulaRule(formula=["COUNTIF($H$5:$H$11,$H5)>1"], fill=MERGE))
widths = [11,10,9,7,9,9,8,9,30,10,7,18,10,10,18,11,11,11]
for j,w in enumerate(widths,1): g.column_dimensions[L(j)].width = w
g.row_dimensions[4].height = 42
g["A15"] = "Как читать"; g["A15"].font = BOLD
notes = ["Недобор > 0 — в лиге меньше людей, чем по цели: зона повышения в неё шире, зона вылета из неё уже.",
 "Таблица № — лиги с одинаковым номером играют в одной сводной таблице; каждый игрок остаётся в своей лиге.",
 "Сводная таблица собирается снизу вверх: лиги копятся, пока не наберётся минимум; остаток сверху прилипает к таблице ниже.",
 "«В каждой группе» — сколько мест в зонах повышения и вылета в одной группе этой таблицы (округление вниз).",
 "«В следующем раунде» — без новичков и оттока; с ними — лист «Прогноз»."]
for i,n in enumerate(notes): g.cell(16+i,1,n).font = GREY
g.freeze_panes = "B5"

# ---------- Прогноз ----------
p = wb.create_sheet("Прогноз")
p["A1"] = "Прогноз на 12 раундов"; p["A1"].font = f(bold=True, size=14)
p["A2"] = "Раунд 1 = ввод. Дальше: перемещения по зонам → отток → новички по калибровке. Числа ожидаемые (дробные), показаны округлёнными."; p["A2"].font = GREY
p["A4"] = "Раунд"
for i,lg in enumerate(leagues):
    p.cell(4,2+i,lg); p.cell(4,12+i,f"Недобор {lg}")
for i in range(7): p.cell(4,21+i,f"Вверх % {leagues[i]}")
for i in range(1,8): p.cell(4,28+i,f"Вниз % {leagues[i]}")
p["J4"] = "Всего"
for c in p[4]:
    if c.value: c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
NC = lambda i: L(2+i); GC = lambda i: L(12+i); UC = lambda i: L(21+i); DC = lambda i: L(28+i)
for rr in range(12):
    r = 5+rr
    p[f"A{r}"] = rr+1
    for i in range(8):
        if rr == 0:
            p[f"{NC(i)}{r}"] = f"='Ввод'!B{5+i}"; p[f"{NC(i)}{r}"].font = GREEN
        else:
            q = r-1
            expr = f"{NC(i)}{q}"
            if i < 7: expr += f"-{NC(i)}{q}*{UC(i)}{q}"
            if i > 0: expr += f"-{NC(i)}{q}*{DC(i)}{q}"
            if i > 0: expr += f"+{NC(i-1)}{q}*{UC(i-1)}{q}"
            if i < 7: expr += f"+{NC(i+1)}{q}*{DC(i+1)}{q}"
            p[f"{NC(i)}{r}"] = f"=({expr})*(1-{CHURN})+{NEWC}*'Ввод'!$D${5+i}"
        tgt = f"'Ввод'!$C${5+i}"
        p[f"{GC(i)}{r}"] = f"=IF({tgt}=0,0,({tgt}-IF($J{r}=0,0,{NC(i)}{r}/$J{r}))/{tgt})"
    p[f"J{r}"] = f"=SUM(B{r}:I{r})"
    for i in range(7):
        p[f"{UC(i)}{r}"] = f"={clamp(f'{BASE}+{K}*{GC(i+1)}{r}')}"
    for i in range(1,8):
        z = clamp(f"{BASE}-{K}*{GC(i)}{r}")
        p[f"{DC(i)}{r}"] = f"=IF({NC(7)}{r}<{IMM},0,{z})" if i == 7 else f"={z}"
    for c in range(1,36):
        cell = p.cell(r,c)
        if cell.font != GREEN: cell.font = NORM
        cell.number_format = "0" if c <= 10 else "0%"
for c in range(12,36):
    for r in range(4,17): p.cell(r,c).font = GREY if r>4 else f(bold=True, color="666666")
p["L3"] = "Служебные колонки: недобор и зоны каждого раунда"; p["L3"].font = GREY
for j in range(1,11): p.column_dimensions[L(j)].width = 9.5
for j in range(12,36): p.column_dimensions[L(j)].width = 10
p.row_dimensions[4].height = 30
p["A19"] = "Раунд 12: факт против цели"; p["A19"].font = BOLD
for j,h in enumerate(["Лига","Игроков","Доля","Цель","Отклонение"],1):
    c = p.cell(20,j,h); c.font = BOLD; c.fill = HEAD
for i in range(8):
    r = 21+i
    p[f"A{r}"] = leagues[i]
    p[f"B{r}"] = f"={NC(i)}16"; p[f"B{r}"].number_format = "0"
    p[f"C{r}"] = f"=IF($J$16=0,0,B{r}/$J$16)"; p[f"C{r}"].number_format = "0.0%"
    p[f"D{r}"] = f"='Ввод'!C{5+i}"; p[f"D{r}"].number_format = "0%"
    p[f"E{r}"] = f"=C{r}-D{r}"; p[f"E{r}"].number_format = '+0.0%;-0.0%;0.0%'
    for col in "ABCDE": p[f"{col}{r}"].font = NORM
# chart: 100% stacked columns of league sizes per round
ch = BarChart(); ch.type = "col"; ch.grouping = "percentStacked"; ch.overlap = 100; ch.gapWidth = 40
ch.title = "Распределение по лигам, раунды 1–12"
ch.y_axis.title = "Доля игроков"; ch.x_axis.title = "Раунд"; ch.y_axis.number_format = "0%"
ch.y_axis.majorGridlines = None
data = Reference(p, min_col=2, max_col=9, min_row=4, max_row=16)
ch.add_data(data, titles_from_data=True)
ch.set_categories(Reference(p, min_col=1, min_row=5, max_row=16))
ramp = ["86B6EF","6DA7EC","5598E7","3987E5","2A78D6","1C5CAB","104281","0D366B"]
for sr,hx in zip(ch.series, ramp):
    sr.graphicalProperties = GraphicalProperties(solidFill=hx)
    sr.graphicalProperties.line.solidFill = "FFFFFF"
ch.legend.position = "r"; ch.height = 9; ch.width = 18
ch.x_axis.delete = False; ch.y_axis.delete = False
p.add_chart(ch, "G19")
p.freeze_panes = "B5"
wb.save("/tmp/claude-0/-home-claude/3be3caeb-c544-53c8-b739-255c4cef4401/scratchpad/leagues/dota-fantasy-ligi-model.xlsx")
